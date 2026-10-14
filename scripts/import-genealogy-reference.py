from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = PROJECT_ROOT / "tmp" / "genealogy-reference.xlsx"
OUTPUT_TS = PROJECT_ROOT / "src" / "data" / "genealogyReference.ts"
REPORT = PROJECT_ROOT / "tmp" / "genealogy-reference-comparison.md"
CURRENT_GENEALOGY = PROJECT_ROOT / "src" / "data" / "genealogy.ts"


def cell(row: dict[str, object], key: str) -> str:
    value = row.get(key)
    if value is None:
        return ""
    return str(value).strip()


def split_names(value: str) -> list[str]:
    if not value:
        return []
    parts = re.split(r"[、,，;；/]+", value)
    return [part.strip() for part in parts if part and part.strip()]


def read_reference(path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(path, read_only=False, data_only=True)
    sheet = workbook.worksheets[0]
    headers = [str(value).strip() if value is not None else "" for value in next(sheet.iter_rows(min_row=3, max_row=3, values_only=True))]
    records: list[dict[str, object]] = []

    for values in sheet.iter_rows(min_row=5, values_only=True):
        row = {headers[index]: values[index] for index in range(min(len(headers), len(values))) if headers[index]}
        person_id = cell(row, "person_id")
        name = cell(row, "姓名*")
        if not person_id or not name or person_id.startswith("例"):
            continue

        records.append(
            {
                "id": person_id,
                "name": name,
                "gender": cell(row, "性别") or "未知",
                "birthYear": cell(row, "出生年"),
                "deathYear": cell(row, "去世年"),
                "lifeNote": cell(row, "生卒说明"),
                "alias": cell(row, "字号/别名"),
                "surname": cell(row, "姓氏"),
                "branch": cell(row, "家族/房支"),
                "generation": cell(row, "世代"),
                "generationName": cell(row, "字辈"),
                "fatherName": cell(row, "父亲姓名"),
                "motherName": cell(row, "母亲姓名"),
                "spouseNames": split_names(cell(row, "配偶姓名(可多个)")),
                "childrenNames": split_names(cell(row, "子女姓名(可多个)")),
                "siblings": split_names(cell(row, "兄弟姐妹")),
                "migration": cell(row, "迁徙/籍贯"),
                "summary": cell(row, "人物简介"),
                "events": cell(row, "主要事迹"),
                "excerpt": cell(row, "原文摘录*"),
                "source": cell(row, "来源章节*"),
                "sourceFile": cell(row, "HTML文件/页码/位置"),
                "confidence": cell(row, "可信度*"),
                "status": cell(row, "整理状态"),
                "conflict": cell(row, "冲突/疑点"),
                "category": cell(row, "人物归类*"),
                "lineageFlag": cell(row, "村内世系标识"),
                "lineageRule": cell(row, "世系处理规则"),
                "mergeBasis": cell(row, "合并依据"),
                "priority": cell(row, "核对优先级"),
                "note": cell(row, "分类备注"),
            }
        )
    return records


def read_current_people() -> list[dict[str, object]]:
    text = CURRENT_GENEALOGY.read_text(encoding="utf-8")
    match = re.search(r"export const genealogyPeople: GenealogyPerson\[\] = ([\s\S]*?);\s*$", text)
    if not match:
        return []
    return json.loads(match.group(1))


def normalize_name(name: str) -> str:
    return name.strip()


def current_parent_name(person: dict[str, object], by_id: dict[str, dict[str, object]], key: str) -> str:
    parent_id = person.get(key)
    if not parent_id:
        return ""
    parent = by_id.get(str(parent_id))
    return str(parent.get("name", "")) if parent else ""


def compare(reference: list[dict[str, object]], current: list[dict[str, object]]) -> str:
    by_name: dict[str, list[dict[str, object]]] = {}
    by_id = {str(person.get("id")): person for person in current}
    for person in current:
        by_name.setdefault(normalize_name(str(person.get("name", ""))), []).append(person)

    missing = []
    father_mismatch = []
    mother_mismatch = []
    child_missing = []

    for ref in reference:
        name = normalize_name(str(ref["name"]))
        current_matches = by_name.get(name, [])
        if not current_matches:
            missing.append(ref)
            continue

        ref_father = normalize_name(str(ref.get("fatherName", "")))
        if ref_father:
            if not any(current_parent_name(person, by_id, "fatherId") == ref_father for person in current_matches):
                father_mismatch.append((ref, [current_parent_name(person, by_id, "fatherId") for person in current_matches]))

        ref_mother = normalize_name(str(ref.get("motherName", "")))
        if ref_mother:
            if not any(current_parent_name(person, by_id, "motherId") == ref_mother for person in current_matches):
                mother_mismatch.append((ref, [current_parent_name(person, by_id, "motherId") for person in current_matches]))

        ref_children = [normalize_name(child) for child in ref.get("childrenNames", [])]
        if ref_children:
            current_child_names = set()
            for person in current_matches:
                for child_id in person.get("childrenIds", []):
                    child = by_id.get(str(child_id))
                    if child:
                        current_child_names.add(str(child.get("name", "")))
            missing_children = [child for child in ref_children if child not in current_child_names]
            if missing_children:
                child_missing.append((ref, missing_children))

    lines = [
        "# 族谱参考表比对报告",
        "",
        f"- 参考表记录数：{len(reference)}",
        f"- 当前页面族谱记录数：{len(current)}",
        f"- 当前族谱缺少的参考人物：{len(missing)}",
        f"- 父亲关系不一致/缺失：{len(father_mismatch)}",
        f"- 母亲关系不一致/缺失：{len(mother_mismatch)}",
        f"- 子女关系不一致/缺失：{len(child_missing)}",
        "",
        "## 优先处理建议",
        "",
        "先看高可信、A类村内世系人物；B/C/D/E类只作事件和来源补充，不直接进入严格族谱树。",
        "",
    ]

    def append_table(title: str, rows: list, formatter, limit: int = 40) -> None:
        lines.extend([f"## {title}", ""])
        if not rows:
            lines.extend(["无。", ""])
            return
        for row in rows[:limit]:
            lines.append(formatter(row))
        if len(rows) > limit:
            lines.append(f"...另有 {len(rows) - limit} 条，后续可继续展开。")
        lines.append("")

    append_table(
        "当前族谱缺少的人物（前40）",
        missing,
        lambda ref: f"- {ref['name']}（{ref.get('category', '')}，{ref.get('source', '')}，可信度：{ref.get('confidence', '')}）",
    )
    append_table(
        "父亲关系不一致/缺失（前40）",
        father_mismatch,
        lambda item: f"- {item[0]['name']}：参考父亲={item[0].get('fatherName', '')}；当前父亲={', '.join([x or '空' for x in item[1]])}",
    )
    append_table(
        "母亲关系不一致/缺失（前40）",
        mother_mismatch,
        lambda item: f"- {item[0]['name']}：参考母亲={item[0].get('motherName', '')}；当前母亲={', '.join([x or '空' for x in item[1]])}",
    )
    append_table(
        "子女关系不一致/缺失（前40）",
        child_missing,
        lambda item: f"- {item[0]['name']}：参考子女未出现在当前树={', '.join(item[1])}",
    )

    return "\n".join(lines)


def write_ts(records: list[dict[str, object]]) -> None:
    body = """export type GenealogyReferencePerson = {
  id: string;
  name: string;
  gender: string;
  birthYear: string;
  deathYear: string;
  lifeNote: string;
  alias: string;
  surname: string;
  branch: string;
  generation: string;
  generationName: string;
  fatherName: string;
  motherName: string;
  spouseNames: string[];
  childrenNames: string[];
  siblings: string[];
  migration: string;
  summary: string;
  events: string;
  excerpt: string;
  source: string;
  sourceFile: string;
  confidence: string;
  status: string;
  conflict: string;
  category: string;
  lineageFlag: string;
  lineageRule: string;
  mergeBasis: string;
  priority: string;
  note: string;
};

export const genealogyReferencePeople: GenealogyReferencePerson[] = """
    body += json.dumps(records, ensure_ascii=False, indent=2)
    body += ";\n"
    OUTPUT_TS.write_text(body, encoding="utf-8")


def main() -> None:
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    records = read_reference(path)
    current = read_current_people()
    write_ts(records)
    REPORT.write_text(compare(records, current), encoding="utf-8")
    print(f"已生成 {OUTPUT_TS.relative_to(PROJECT_ROOT)}，共 {len(records)} 条参考记录。")
    print(f"已生成 {REPORT.relative_to(PROJECT_ROOT)}。")


if __name__ == "__main__":
    main()
